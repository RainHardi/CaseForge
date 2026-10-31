"""Turn uploaded files (PDF, Word, Excel, PowerPoint, text) into plain text notes.

The text is shown to the user in the notes box before anything is generated, so they can see
exactly what CaseForge will use — and delete anything confidential.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

TEXT_TYPES = ["txt", "md", "csv", "json", "log"]
DOC_TYPES = ["pdf", "docx", "xlsx", "pptx"]
IMAGE_TYPES = ["png", "jpg", "jpeg"]
ALL_TYPES = TEXT_TYPES + DOC_TYPES + IMAGE_TYPES
MAX_CHARS = 40_000  # keeps one request well inside the model's limits


class IngestError(ValueError):
    pass


def is_image(name: str) -> bool:
    return Path(name).suffix.lower().lstrip(".") in IMAGE_TYPES


def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower().lstrip(".")
    try:
        if ext in ("txt", "md", "log", "json"):
            text = _decode(data)
        elif ext == "csv":
            rows = csv.reader(io.StringIO(_decode(data)))
            text = "\n".join(" | ".join(c.strip() for c in r if c.strip()) for r in rows)
        elif ext == "pdf":
            text = _pdf(data)
        elif ext == "docx":
            text = _docx(data)
        elif ext == "xlsx":
            text = _xlsx(data)
        elif ext == "pptx":
            text = _pptx(data)
        elif ext == "doc":
            raise IngestError("Old .doc files aren't supported — save it as .docx or PDF first.")
        else:
            raise IngestError(f"Can't read .{ext} files.")
    except IngestError:
        raise
    except Exception as e:  # corrupt or password-protected files
        raise IngestError(f"Couldn't read {name}: {e}") from e

    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text)
    if not text.strip():
        if ext == "pdf":
            raise IngestError(f"{name} has no selectable text (probably a scan). Copy the text in by hand.")
        raise IngestError(f"{name} is empty.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n[… file cut here — too long; paste only the relevant part]"
    return text.strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp949", "latin-1"):  # cp949 = Korean Windows files
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise IngestError("This PDF is password-protected.")
    return "\n\n".join((p.extract_text() or "") for p in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat; dedupe
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[Sheet: {ws.title}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"(speaker notes) {notes}")
        if lines:
            parts.append(f"[Slide {i}]\n" + "\n".join(lines))
    return "\n\n".join(parts)
