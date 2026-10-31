"""Uploaded files become plain-text notes."""

import io

import pytest

from caseforge.ingest import IngestError, extract_text, is_image

LINE = "We built a dashboard with Python and React"


def test_txt_and_korean_windows_encoding():
    assert extract_text("notes.txt", LINE.encode()) == LINE
    assert "대시보드" in extract_text("notes.txt", "대시보드를 만들었다".encode("cp949"))


def test_csv():
    assert extract_text("t.csv", b"tool,use\nPython,backend\n") == "tool | use\nPython | backend"


def test_docx_with_table():
    import docx

    d = docx.Document()
    d.add_paragraph(LINE)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Result", "Friday count is faster"
    buf = io.BytesIO(); d.save(buf)
    text = extract_text("report.docx", buf.getvalue())
    assert LINE in text and "Result | Friday count is faster" in text


def test_xlsx():
    from openpyxl import Workbook

    wb = Workbook(); ws = wb.active; ws.title = "Results"
    ws.append(["Week", "No-shows"]); ws.append([1, 18])
    buf = io.BytesIO(); wb.save(buf)
    text = extract_text("data.xlsx", buf.getvalue())
    assert "[Sheet: Results]" in text and "1 | 18" in text


def test_pptx():
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Inventory project"
    s.placeholders[1].text = LINE
    buf = io.BytesIO(); prs.save(buf)
    text = extract_text("deck.pptx", buf.getvalue())
    assert "Inventory project" in text and LINE in text


def test_pdf():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(72, 720, LINE); c.save()
    assert LINE in extract_text("notes.pdf", buf.getvalue())


def test_empty_pdf_gives_a_clear_message():
    from reportlab.pdfgen import canvas

    buf = io.BytesIO(); c = canvas.Canvas(buf); c.showPage(); c.save()
    with pytest.raises(IngestError, match="scan"):
        extract_text("scan.pdf", buf.getvalue())


def test_bad_and_unsupported_files():
    with pytest.raises(IngestError):
        extract_text("broken.docx", b"not a zip")
    with pytest.raises(IngestError, match=".docx"):
        extract_text("old.doc", b"x")
    assert is_image("a.PNG") and not is_image("a.pdf")
